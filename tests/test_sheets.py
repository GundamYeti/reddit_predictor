"""Spreadsheet export and import round trip.

Run: python -m unittest discover tests
"""

import os
import tempfile
import unittest

from predictor import db, sheets

T0 = 1759000000


class Sheets(unittest.TestCase):
    def setUp(self):
        self.conn = db.connect(":memory:")
        for cid, body in [("p1", "He'll drop out"), ("p2", "He should drop out")]:
            self.conn.execute("INSERT INTO comments (id, author, subreddit, body, created_utc, source)"
                              " VALUES (?, 'u', 'politics', ?, ?, 'pilot1')", (cid, body, T0))
        self.conn.execute("INSERT INTO comments (id, author, subreddit, body, created_utc, source)"
                          " VALUES ('k1', 'kim', 'nba', '', ?, 'panel')", (T0,))
        self.conn.execute("INSERT INTO forecasts (comment_id, extractor, proposition, direction, hedge_phrase)"
                          " VALUES ('k1', 'm', 'Knicks win the title', 'happens', 'probably')")
        for q in ("Will the Knicks win the 2027 title?", "Will the Jazz make the playoffs?"):
            cur = self.conn.execute("INSERT INTO claims (question, tag) VALUES (?, 'panel')", (q,))
            self.conn.execute("INSERT INTO claim_members (claim_id, forecast_id, stance) VALUES (?, 1, 'yes')",
                              (cur.lastrowid,))

    def test_round_trip(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "review.xlsx")
            counts = sheets.export(self.conn, path)
            self.assertEqual(counts, {"labels": 2, "claims": 2})
            wb = load_workbook(path)
            # The Label tab is blind: no model answers in it.
            self.assertNotIn("Knicks", " ".join(str(c.value) for row in wb["Label"].iter_rows() for c in row))
            label = wb["Label"]
            for row in label.iter_rows(min_row=2):
                row[4].value = "F" if row[0].value == "p1" else "n"
            claims_ws = wb["Claims"]
            for row in claims_ws.iter_rows(min_row=2):
                row[10].value = "y"
                if "Knicks" in row[1].value:
                    row[11].value = "2027-06-30"      # the other has no date and no market: a problem
            wb.save(path)
            log = []
            saved, created, problems = sheets.import_(self.conn, path, out=log.append)
        self.assertEqual((saved, created), (2, 1))
        self.assertEqual(dict(self.conn.execute("SELECT comment_id, label FROM labels")),
                         {"p1": "forecast", "p2": "not"})
        self.assertEqual(len(problems), 1)
        self.assertIn("no matched market", problems[0])
        q = self.conn.execute("SELECT * FROM questions").fetchone()
        self.assertEqual(q["text"], "Will the Knicks win the 2027 title?")


if __name__ == "__main__":
    unittest.main()
