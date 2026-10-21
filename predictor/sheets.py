"""Do the hand work in a spreadsheet instead of the terminal, on any computer.

  export: one .xlsx with
    - Label:  the pilot's labeling set (blind: no model answers shown)
    - Claims: candidate questions, with a column to approve each
  import: read your answers back: labels go to the labels table, approved
          claims are promoted to questions (as `claims promote` would)

Only the yellow columns are read back; everything else can be sorted or
filtered freely, but don't delete the id columns.
"""

import datetime as dt
import time

from . import claims, pilot

LABELS = {"f": "forecast", "c": "conditional", "n": "not", "u": "unsure"}
FONT = "Arial"


def _styles():
    from openpyxl.styles import Alignment, Font, PatternFill
    return {
        "head": Font(name=FONT, bold=True, color="FFFFFF"),
        "head_input": Font(name=FONT, bold=True, color="000000"),
        "title": Font(name=FONT, bold=True, size=12),
        "head_fill": PatternFill("solid", fgColor="44546A"),
        "input_fill": PatternFill("solid", fgColor="FFFF00"),
        "body": Font(name=FONT, size=10),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "top": Alignment(vertical="top"),
    }


def _table(ws, headers, widths, rows, inputs, st):
    ws.append(headers)
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=i)
        cell.font = st["head_input"] if h in inputs else st["head"]
        cell.fill = st["input_fill"] if h in inputs else st["head_fill"]
        ws.column_dimensions[cell.column_letter].width = w
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = st["body"]
            cell.alignment = st["wrap"]
            if headers[cell.column - 1] in inputs:
                cell.fill = st["input_fill"]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def export(conn, path, tag="pilot1", claim_tag=None):
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    st = _styles()
    wb = Workbook()

    # Instructions first, with the legend and an example of each.
    ws = wb.active
    ws.title = "Instructions"
    lines = [
        ("How to use this workbook", True),
        ("Fill in only the YELLOW columns. Save the file and send it back; "
         "`python -m predictor sheets import <file>` reads your answers.", False),
        ("", False),
        ("Label tab: is each comment a forecast?", True),
        ("Your label: f = a forecast (the author expects something checkable to happen or not), "
         "c = a conditional forecast (\"if A, then B\"), n = not a forecast (wish, opinion, question, "
         "joke, the past), u = can't tell. Note is optional.", False),
        ("Example:  \"Calling it now, Harris is the nominee within a month.\"  ->  f", False),
        ("Example:  \"He needs to drop out.\"  ->  n   (a wish, not an expectation)", False),
        ("", False),
        ("Claims tab: which claims should become questions?", True),
        ("Approve: y to make it a question. Claims with a Polymarket market resolve with the market's "
         "rules and outcome; for the others, also give Closes (YYYY-MM-DD, when the answer is known) "
         "and, if you like, Your resolution rule.", False),
        ("Example:  approve = y, closes = 2026-11-04, rule = \"YES if Abdul El-Sayed is certified "
         "the winner of the Michigan Senate general election\"", False),
        ("Reword (optional) replaces the question text.", False),
        ("", False),
        (f"Exported {time.strftime('%Y-%m-%d %H:%M')} from data/predictor.db.", False),
    ]
    for text, bold in lines:
        ws.append([text])
        c = ws.cell(row=ws.max_row, column=1)
        c.font = st["title"] if bold else st["body"]
        c.alignment = st["wrap"]
    ws.column_dimensions["A"].width = 110

    # Label: the pilot labeling set, minus anything already labeled.
    ws = wb.create_sheet("Label")
    done = {r[0] for r in conn.execute("SELECT comment_id FROM labels")}
    rows = []
    for c in pilot.label_set(conn, tag):
        if c["id"] in done:
            continue
        rows.append([c["id"], "r/" + c["subreddit"], time.strftime("%Y-%m-%d", time.gmtime(c["created_utc"])),
                     c["body"][:3000], None, None])
    _table(ws, ["Comment id", "Subreddit", "Posted", "Comment", "Your label", "Note"],
           [12, 14, 11, 100, 11, 30], rows, {"Your label", "Note"}, st)
    dv = DataValidation(type="list", formula1='"f,c,n,u"', allow_blank=True,
                        error="Use f, c, n or u", errorTitle="Label")
    ws.add_data_validation(dv)
    dv.add(f"E2:E{max(2, ws.max_row)}")

    # Claims: everything not yet promoted, most-predicted first.
    ws = wb.create_sheet("Claims")
    rows = []
    for r in claims.listing(conn, claim_tag, min_authors=1, top=100000):
        if r["question_id"]:
            continue
        if r["market_question"]:
            status = (f"resolved {r['outcome'].upper()}" if r["outcome"]
                      else f"{r['price_yes']:.0%} yes now" if r["price_yes"] is not None else "open")
        else:
            status = ""
        rows.append([r["id"], r["question"], r["deadline"], r["authors"], r["forecasts"], r["yes"], r["no"],
                     ", ".join("r/" + s for s in r["subreddits"].split(",")), r["market_question"] or "", status,
                     None, None, None, None])
    _table(ws, ["Claim id", "Question", "Time frame", "People", "Forecasts", "Yes", "No", "Subreddits",
                "Polymarket market", "Market status", "Approve", "Closes", "Your resolution rule", "Reword"],
           [9, 60, 16, 8, 10, 6, 6, 22, 45, 16, 9, 12, 40, 30], rows,
           {"Approve", "Closes", "Your resolution rule", "Reword"}, st)
    dv = DataValidation(type="list", formula1='"y"', allow_blank=True, error="Use y, or leave blank",
                        errorTitle="Approve")
    ws.add_data_validation(dv)
    dv.add(f"K2:K{max(2, ws.max_row)}")

    wb.save(path)
    return {"labels": wb["Label"].max_row - 1, "claims": wb["Claims"].max_row - 1}


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        d = v if isinstance(v, dt.datetime) else dt.datetime(v.year, v.month, v.day)
    else:
        d = dt.datetime.strptime(str(v).strip()[:10], "%Y-%m-%d")
    return int(d.replace(tzinfo=dt.timezone.utc).timestamp())


def import_(conn, path, out=print):
    """Read labels and approvals back. Returns (labels saved, questions created, problems)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    saved = created = 0
    problems = []
    if "Label" in wb.sheetnames:
        ws = wb["Label"]
        head = [c.value for c in ws[1]]
        ci, li, ni = head.index("Comment id"), head.index("Your label"), head.index("Note")
        for row in ws.iter_rows(min_row=2, values_only=True):
            label = str(row[li] or "").strip().lower()[:1]
            if not row[ci] or not label:
                continue
            if label not in LABELS:
                problems.append(f"Label tab, comment {row[ci]}: {row[li]!r} isn't f, c, n or u")
                continue
            conn.execute("INSERT OR REPLACE INTO labels (comment_id, label, note) VALUES (?, ?, ?)",
                         (str(row[ci]), LABELS[label], str(row[ni] or "")[:500]))
            saved += 1
        conn.commit()
    if "Claims" in wb.sheetnames:
        ws = wb["Claims"]
        head = [c.value for c in ws[1]]
        idx = {h: head.index(h) for h in ("Claim id", "Approve", "Closes", "Your resolution rule", "Reword")}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[idx["Claim id"]] or str(row[idx["Approve"]] or "").strip().lower() != "y":
                continue
            cid = int(row[idx["Claim id"]])
            try:
                closes = _date(row[idx["Closes"]])
            except ValueError:
                problems.append(f"Claims tab, claim {cid}: Closes {row[idx['Closes']]!r} isn't a YYYY-MM-DD date")
                continue
            try:
                qid = claims.promote(conn, cid, closes, row[idx["Reword"]] or None,
                                     row[idx["Your resolution rule"]] or None)
                created += 1
                out(f"  claim #{cid} -> question #{qid}")
            except ValueError as e:
                problems.append(f"Claims tab, claim {cid}: {e}")
    out(f"Saved {saved} labels and created {created} questions.")
    for p in problems:
        out("  problem: " + p)
    return saved, created, problems
